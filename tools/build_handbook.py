"""根据 data/points.js 生成《五角场骑手地图调研手册》Word 稿（可作为成果报告册）。

用法：python tools/build_handbook.py [输出路径]
默认输出到仓库上一级目录：
    真实数据 → ../调研手册（自动生成稿）.docx
    示例数据 → ../调研手册（示例数据预览）.docx

点位图鉴、统计表、采集概况由数据自动生成；有 docs/反馈问卷.md 时自动统计骑手反馈；
有 docs/screenshots/ 和 docs/海报.png 时自动插入网页截图和海报。
黄色高亮的【待填写】处需要调研组补充。
每次重新生成会覆盖旧文件，所以请在生成后的副本上修改，或先改完数据再生成。
"""
import re
import sys
from collections import Counter
from pathlib import Path

from docx.shared import Cm, RGBColor

from common import PAGES_URL, PHOTOS_DIR, REPO_URL, ROOT, load_points, load_taxonomy
from docx_util import (GRAY, GREEN, bullet, kv_table, new_document, page_break, page_numbers, para, rich,
                       steps, table, todo)
from feedback import load_responses, tally, texts
from make_guide import INTERVIEW_OPENING, INTERVIEW_QUESTIONS

SHOTS = ROOT / "docs" / "screenshots"
POSTER = ROOT / "docs" / "海报.png"


# 封面信息（来自申报书）
PROJECT = "技术向善——杨浦区外卖配送调研与信息地图实践"
SCHOOL = "复旦大学计算与智能创新学院"
LEADER = "赵思远"
ADVISOR = "蔡瑞悦"
MEMBERS = "李昀泽、唐子姗、刘沛瑜、徐紫羽、张蕙孜"
DATE = "2026 年 12 月"
RED = RGBColor(0xC0, 0x2A, 0x1A)


def short(option):
    """去掉选项里较长的括号说明（如“（能不能喝水、充电、上厕所）”），表格里更紧凑；
    “（专送）”“（众包）”这类区分选项的短括号保留。"""
    return re.sub(r"（[^）]*[、，][^）]*）", "", option).strip()


def picture(doc, paths, width_cm, caption):
    """一行放一张或多张图，下面加图注。找不到图片时返回 False。"""
    paths = [p for p in paths if p.exists()]
    if not paths:
        return False
    par = doc.add_paragraph()
    par.alignment = 1  # 居中
    for i, f in enumerate(paths):
        if i:
            par.add_run("　")
        par.add_run().add_picture(str(f), width=Cm(width_cm))
    para(doc, caption, size=9.5, color=GRAY, align="center")
    return True


def feedback_section(doc, num):
    qs, rs = load_responses()
    doc.add_heading(f"{num}、骑手反馈", level=1)
    if not rs:
        todo(doc, "回访后把答卷按格式整理进 docs/反馈问卷.md，重新生成本手册即可自动统计。")
        return
    def is_rider(r):
        return not any(c.startswith("不是骑手") for c in r["answers"].get(1, {}).get("choices", []))
    riders = [r for r in rs if is_rider(r)]
    para(doc, f"共收回反馈问卷 {len(rs)} 份，其中骑手 {len(riders)} 份、非骑手 {len(rs) - len(riders)} 份。"
              "样本较少，以下只描述答卷情况，不代表五角场骑手整体。")
    todo(doc, "写明问卷在什么时间、什么场合收集（当面回访还是线上），以及是本人填写还是调研员代填。")

    doc.add_heading("（一）选择题统计", level=2)
    rows = [[f"{q}. {qs[q]}", n, "；".join(f"{short(k)} {v}" for k, v in c.most_common())]
            for q, (n, c) in tally(qs, rs).items()]
    table(doc, ["题目", "作答人数", "选项（人数）"], rows, [6, 1.6, 8.4], font_size=9.5)

    if riders:
        doc.add_heading("（二）每位骑手的回答", level=2)
        def pick(r, q):
            a = r["answers"].get(q)
            return "、".join(short(x) for x in a["choices"]) if a and a["choices"] else "—"
        rows = [[r["title"], pick(r, 3), pick(r, 6), pick(r, 7), pick(r, 8)] for r in riders]
        table(doc, ["受访者", "之前知道歇脚点吗", "觉得有用的", "不方便的", "还会打开吗"], rows,
              [3.6, 2.6, 3.8, 3.6, 2.4], font_size=9.5)

    doc.add_heading("（三）原话摘录", level=2)
    for q, items in texts(qs, rs).items():
        items = [(t, x) for t, x in items if x.strip() not in ("没有", "无", "没")]
        if not items:
            continue
        para(doc, f"{q}. {qs[q]}", bold=True, size=10.5)
        for title, x in items:
            bullet(doc, f"{x}（{title}）")
    para(doc, "答卷中提到的地点已整理进 docs/待核实线索.xlsx，调研组实地核实后才会加入地图。", size=10, color=GRAY)

    doc.add_heading("（四）解读", level=2)
    todo(doc, "结合上面的表格写 2–3 句解读。例如：新骑手和老骑手各自更需要哪类信息；不再打开的原因是什么；对网页和对平台分别意味着什么。")


def point_card(doc, p, sub_label, cat_label):
    doc.add_heading(p["name"], level=3)
    pairs = [("类型", f"{cat_label} · {sub_label}"), ("地址", p.get("address", ""))]
    if p.get("desc"):
        pairs.append(("说明", p["desc"]))
    if p.get("facilities"):
        pairs.append(("设施", p["facilities"]))
    if p.get("hours"):
        pairs.append(("开放时间", p["hours"]))
    if p.get("tags"):
        pairs.append(("标签", "、".join(p["tags"])))
    if p.get("quote"):
        pairs.append(("骑手原话", "“" + p["quote"].strip("“”\"") + "”"))
    pairs.append(("来源", f'{p.get("source", "")}　{p.get("date", "")}　{p.get("collector", "")}'.strip()))
    if p.get("refs"):
        pairs.append(("出处", "\n".join(p["refs"])))
    pairs.append(("坐标", f'{p["lng"]:.6f}, {p["lat"]:.6f}（GCJ-02）'))
    if p.get("locNote"):
        pairs.append(("位置说明", p["locNote"]))
    kv_table(doc, pairs)

    photos = [PHOTOS_DIR / f for f in p.get("photos", [])[:2]]
    photos = [f for f in photos if f.exists()]
    if photos:
        para_ = doc.add_paragraph()
        for f in photos:
            para_.add_run().add_picture(str(f), width=Cm(7.6))
            para_.add_run("  ")
    doc.add_paragraph()


def main():
    tax = load_taxonomy()
    is_sample, generated_at, points = load_points()
    default_name = "调研手册（示例数据预览）.docx" if is_sample else "调研手册（自动生成稿）.docx"
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / default_name

    sub_info = {s["key"]: (c, s) for c in tax["categories"] for s in c["subtypes"]}
    by_cat = {c["key"]: [p for p in points if p["category"] == c["key"]] for c in tax["categories"]}
    sub_count = Counter(p["subtype"] for p in points)
    src_count = Counter(p.get("source", "") for p in points)

    doc = new_document()
    page_numbers(doc)

    # ---------- 封面 ----------
    for _ in range(5):
        doc.add_paragraph()
    para(doc, "五角场骑手地图", bold=True, size=30, align="center")
    para(doc, "调 研 手 册", bold=True, size=20, color=GREEN, align="center")
    doc.add_paragraph()
    para(doc, PROJECT, size=13, align="center")
    para(doc, "2026 年秋季日常化社会实践 · 人民城市建设专题", size=11, color=GRAY, align="center")
    for _ in range(6):
        doc.add_paragraph()
    for line in (SCHOOL, f"项目负责人：{LEADER}　　指导老师：{ADVISOR}", f"项目成员：{MEMBERS}", DATE):
        para(doc, line, size=12, align="center")
    if is_sample:
        doc.add_paragraph()
        para(doc, "注意：本稿由示例数据生成，点位均为虚构，仅用于预览排版。拿到真实点位后请重新生成。",
             bold=True, color=RED, align="center")
    page_break(doc)

    # ---------- 一、调研背景 ----------
    doc.add_heading("一、调研背景", level=1)
    para(doc, "五角场位于上海市杨浦区，万达广场、合生汇、百联又一城等大型商场集中，一侧连着创智天地园区，"
              "一侧连着复旦大学、上海财经大学等高校，商场、写字楼、高校和居民区叠加，外卖订单十分密集。")
    para(doc, "2026 年春季，团队曾在五角场商圈访谈外卖骑手，发现两个突出问题：")
    bullet(doc, "受访骑手普遍认为平台给定的送餐时间不够用。电动车进不了的校园、要负重步行很远的商场写字楼，"
                "都没有被算进时间里。", "算法看不见的障碍。")
    bullet(doc, "五角场街道“新心驿站”饮水、充电、微波炉等设施齐全，但骑手反映“位置太远”“怕错过订单”"
                "“不知道有这个地方”，更习惯在路边休息。", "驿站“看得见、用不上”。")
    para(doc, "本学期，团队把实地走访、骑手访谈和公开资料中找到的歇脚点和障碍点整理成本手册，"
              "并做成可在手机上直接打开的《五角场骑手地图》网页，让骑手能找到歇脚点、提前知道哪里难送。")
    todo(doc, "可补充：团队为什么关注这个话题、和“人民城市建设”主题的联系（2–3 句）。")

    # ---------- 二、调研方法 ----------
    doc.add_heading("二、调研方法", level=1)
    has = lambda name: any(name in (p.get("source") or "") for p in points)
    methods = []
    if has("实地走访") or is_sample:
        methods.append(("实地走访：", "队员到现场记录设施和使用情况，用网页“坐标拾取”功能取坐标、拍照、写说明。"))
    if has("骑手指认") or is_sample:
        methods.append(("骑手访谈：", "利用骑手等餐、休息的间隙进行简短访谈，请骑手指认附近的歇脚点和障碍点（提纲见附录一）。"))
    if has("公开资料"):
        methods.append(("公开资料检索：", "检索杨浦区政府网站和新闻报道中公布的骑手驿站、商圈治理等信息，逐条记录出处链接。"
                                        "这类点位没有经过实地核实，在地图和本手册中单独标注。"))
    if load_responses()[1]:
        methods.append(("反馈问卷：", "网页完成后请骑手试用并填写反馈问卷，统计见第六章。"))
    methods.append(("数据整理：", "统一填写点位采集表，由脚本校验必填项、分类、坐标范围和出处，并去除照片中的 GPS 等拍摄信息。"))
    steps(doc, methods)
    if points and not is_sample:
        located = sum(1 for p in points if p.get("locNote"))
        para(doc, f"地图共收录点位 {len(points)} 个，按来源分：" +
             "，".join(f"{k} {v} 个" for k, v in src_count.most_common() if k) + "。" +
             (f"其中 {located} 个点位的坐标按地图标注标定，具体见各点位的“位置说明”。" if located else ""))
    todo(doc, "写明线下活动的日期、地点和参与人员，以及访谈了几位骑手（平台、从业时长等，不写姓名）。")
    para(doc, "隐私保护：访谈前说明用途；手册、网页和代码仓库均不包含骑手姓名、正脸照片、手机号等个人信息，骑手原话均匿名。",
         size=10, color=GRAY)

    # ---------- 三、点位总览 ----------
    doc.add_heading("三、点位总览", level=1)
    rows = []
    for c in tax["categories"]:
        for s in c["subtypes"]:
            rows.append([c["label"], f'{s["icon"]} {s["label"]}', sub_count.get(s["key"], 0), s["hint"]])
    rows.append(["合计", "", len(points), ""])
    table(doc, ["大类", "子类", "数量", "包括哪些地方"], rows, [2, 3.4, 1.6, 9])
    if not picture(doc, [SHOTS / "desktop.png"], 16, "图 1　《五角场骑手地图》网页（电脑端）"):
        todo(doc, "插入网页地图截图（运行 tools/take_screenshots.py 后重新生成即可自动插入）。")

    # ---------- 四、五、图鉴 ----------
    for num, c in (("四", tax["categories"][0]), ("五", tax["categories"][1])):
        page_break(doc)
        doc.add_heading(f"{num}、{c['label']}图鉴", level=1)
        pts = by_cat.get(c["key"], [])
        if not pts:
            todo(doc, f"暂无{c['label']}数据。")
            continue
        for s in c["subtypes"]:
            group = [p for p in pts if p["subtype"] == s["key"]]
            if not group:
                continue
            doc.add_heading(f'{s["icon"]} {s["label"]}（{len(group)}）', level=2)
            for p in sorted(group, key=lambda x: x["name"]):
                point_card(doc, p, s["label"], c["label"])

    # ---------- 六、骑手反馈 ----------
    page_break(doc)
    feedback_section(doc, "六")

    # ---------- 七、发现与建议 ----------
    page_break(doc)
    doc.add_heading("七、发现与建议", level=1)
    if points:
        obs = [(k, v) for k, v in sub_count.most_common() if sub_info[k][0]["key"] == "obstacle"]
        # 只有一类明显最多时才写“最多的是”，并列时不写
        top_obstacle = sub_info[obs[0][0]][1]["label"] if obs and (len(obs) == 1 or obs[0][1] > obs[1][1]) else None
        facts = f"从数据看：共收录歇脚点 {len(by_cat.get('rest', []))} 个、障碍点 {len(by_cat.get('obstacle', []))} 个"
        if top_obstacle:
            facts += f"；障碍点中最多的是“{top_obstacle}”"
        para(doc, facts + "。")
    doc.add_heading("（一）发现", level=2)
    todo(doc, "结合点位和访谈写 2–3 条发现。可参考的方向：歇脚点集中在哪里、离订单密集区有多远；"
              "障碍点集中在哪些学校、商场、小区；骑手怎么评价驿站；和春季调研相比有什么变化。")
    doc.add_heading("（二）建议", level=2)
    para(doc, "以下为春季调研提出的方向，请结合本次数据修改或补充：", size=10, color=GRAY)
    bullet(doc, "把高校、禁行楼宇、大型商场等障碍标注成标签，送餐时间按场景给出合理补偿，而不是统一加几分钟。",
           "给平台：")
    bullet(doc, "在骑手常用的渠道推送附近驿站的位置和设施；与高频出餐商户合作设立小型服务点，把服务“送到骑手手边”。",
           "给驿站和街道：")
    bullet(doc, "持续更新这张地图，欢迎骑手和市民反馈新的点位。", "给我们自己：")
    if points and not is_sample:
        doc.add_heading("（三）局限", level=2)
        limits = [f"点位数量较少（{len(points)} 个），只覆盖五角场一带的一部分，不能代表全部情况。"]
        if has("公开资料"):
            limits.append("部分点位来自公开资料，没有实地核实，开放时间、设施等可能已有变化。")
        if any(p.get("locNote") for p in points):
            limits.append("部分坐标是按地图标注标定的，不是现场所取，可能有几十米的偏差。")
        if load_responses()[1]:
            limits.append("反馈问卷样本很少，只能说明个别骑手的看法。")
        for x in limits:
            bullet(doc, x)

    # ---------- 八、网页介绍 ----------
    doc.add_heading("八、《五角场骑手地图》网页", level=1)
    para(doc, "网页用手机浏览器即可打开，不需要下载 App。主要功能：")
    for name, text in (
        ("地图浏览：", "缩放、拖动，一键回到五角场；歇脚点为绿色、障碍点为橙红色。"),
        ("类型筛选：", "按大类和子类勾选显示。"),
        ("搜索：", "按名称、地址、标签搜索；搜“厕所”也能找到“卫生间”；搜索框下有喝水、厕所、充电等常用按钮，点一下就搜。"),
        ("详情：", "说明、设施、开放时间、照片、骑手原话和信息来源。"),
        ("我的位置与导航：", "定位后按距离排序，一键跳转高德地图导航。"),
        ("大字模式：", "右上角“大字”按钮放大全部文字，方便在路边看清。"),
        ("坐标拾取：", "调研时在现场一键获取坐标，方便以后继续补充点位。"),
    ):
        bullet(doc, text, name)
    para(doc, "其中“大字模式”“常用搜索按钮”和近义词搜索，是根据反馈问卷里骑手提出的“字太小”“找不到想要的地方”新增的。",
         size=10, color=GRAY)
    picture(doc, [SHOTS / "mobile.png", SHOTS / "mobile-detail.png"], 6.2, "图 2　手机端：地图与列表（左），点位详情（右）")
    rich(doc, ["网页地址：", (PAGES_URL, {"bold": True})])
    rich(doc, ["代码开源：", (REPO_URL, {"bold": True})])
    if POSTER.exists():
        para(doc, "宣传海报（带二维码，可贴在驿站和商家门口）见附录三。")

    # ---------- 附录 ----------
    page_break(doc)
    doc.add_heading("附录一　骑手访谈提纲", level=1)
    para(doc, "开场白：", bold=True)
    para(doc, INTERVIEW_OPENING)
    steps(doc, INTERVIEW_QUESTIONS)
    doc.add_heading("附录二　活动照片", level=1)
    todo(doc, "插入活动照片（8 张，命名为“拍摄者+照片内容简介”）。")
    if POSTER.exists():
        page_break(doc)
        doc.add_heading("附录三　宣传海报", level=1)
        picture(doc, [POSTER], 12.5, "扫码打开《五角场骑手地图》")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"已生成 {out}")
    if is_sample:
        print("（当前是示例数据，仅用于预览排版）")


if __name__ == "__main__":
    main()
