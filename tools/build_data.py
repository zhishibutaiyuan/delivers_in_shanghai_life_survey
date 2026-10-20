"""把调研组的点位表（data/points.xlsx）转换成网页数据（data/points.js），并处理照片。

用法：
    python tools/build_data.py            # 校验 + 生成 data/points.js + 处理照片
    python tools/build_data.py --check    # 只校验，不写文件
    python tools/build_data.py 某个表.xlsx  # 指定其他表格

校验不通过时不会覆盖 data/points.js，会逐行列出问题。
照片：从 data/photos_raw/ 读取原图，旋转摆正、缩小到长边 1600 像素、去除 GPS 等拍摄信息，
保存到 photos/。
"""
import argparse
import datetime as dt
import hashlib
import json
import re
import sys

from openpyxl import load_workbook

from common import (COLUMNS, PHOTOS_DIR, POINTS_JS, POINTS_XLSX, RAW_PHOTOS_DIR, ROOT, YANGPU_BBOX,
                    load_taxonomy, split_list)

PHOTO_MAX_SIDE = 1600
PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_date(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = cell_text(v).replace("/", "-").replace(".", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if not m:
        m2 = re.match(r"^(\d{4})年(\d{1,2})月(\d{1,2})日?$", s)
        m = m2
    if not m:
        return None
    y, mo, d = (int(x) for x in m.groups())
    try:
        return dt.date(y, mo, d).strftime("%Y-%m-%d")
    except ValueError:
        return None


def read_rows(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["点位"] if "点位" in wb.sheetnames else wb.worksheets[0]
    header = [cell_text(c.value).rstrip(" *") for c in ws[1]]
    missing = [h for h, *_ in COLUMNS if h not in header]
    if missing:
        raise SystemExit(f"表头缺少这些列：{'、'.join(missing)}。请使用 docs/点位采集表.xlsx 模板。")
    idx = {h: header.index(h) for h, *_ in COLUMNS}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {h: (row[i] if i < len(row) else None) for h, i in idx.items()}
        if all(cell_text(v) == "" for v in values.values()):
            continue
        yield r, values


def make_id(name, subtype):
    return hashlib.sha1(f"{subtype}|{name}".encode("utf-8")).hexdigest()[:8]


def process_photo(name, errors_out, do_write):
    """返回处理后的文件名；出错时写入 errors_out 并返回 None。"""
    raw = RAW_PHOTOS_DIR / name
    done_name = re.sub(r"\.(jpe?g|png|heic|webp)$", "", name, flags=re.I) + ".jpg"
    done = PHOTOS_DIR / done_name
    if not raw.exists():
        if done.exists():
            return done_name  # 之前已处理过
        errors_out.append(f"找不到照片 {name}（应放在 data/photos_raw/ 里）")
        return None
    if not do_write:
        return done_name
    try:
        from PIL import Image, ImageOps
    except ImportError:
        raise SystemExit("处理照片需要 Pillow：pip install pillow")
    try:
        with Image.open(raw) as im:
            im = ImageOps.exif_transpose(im)  # 按拍摄方向摆正
            im = im.convert("RGB")
            im.thumbnail((PHOTO_MAX_SIDE, PHOTO_MAX_SIDE))
            PHOTOS_DIR.mkdir(exist_ok=True)
            # 不传 exif 参数，保存时就不带任何拍摄信息（含 GPS）
            im.save(done, "JPEG", quality=82, optimize=True, progressive=True)
    except Exception as e:  # noqa: BLE001
        errors_out.append(f"照片 {name} 无法处理：{e}")
        return None
    return done_name


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="?", default=str(POINTS_XLSX))
    ap.add_argument("--check", action="store_true", help="只校验，不写文件")
    args = ap.parse_args()

    tax = load_taxonomy()
    cat_by_label = {c["label"]: c for c in tax["categories"]}
    sub_by_label = {}
    for c in tax["categories"]:
        for s in c["subtypes"]:
            sub_by_label[s["label"]] = (c["key"], s["key"])
    sources = {s: s for s in tax["sources"]}
    sources.update({"两者都有": "实地走访+骑手指认", "实地走访＋骑手指认": "实地走访+骑手指认"})

    problems = []   # (行号, 问题)
    warnings = []
    points = []
    seen = {}

    for r, v in read_rows(args.xlsx):
        errs = []
        t = {h: cell_text(v[h]) for h, *_ in COLUMNS}

        for head, _, required, _ in COLUMNS:
            if required and head not in ("经度", "纬度") and not t[head]:
                errs.append(f"“{head}”没填")

        cat = cat_by_label.get(t["大类"])
        if t["大类"] and not cat:
            errs.append(f"大类“{t['大类']}”不对，只能是 歇脚点 / 障碍点")
        sub = sub_by_label.get(t["子类"])
        if t["子类"] and not sub:
            errs.append(f"子类“{t['子类']}”不在分类表里")
        if cat and sub and sub[0] != cat["key"]:
            errs.append(f"子类“{t['子类']}”不属于大类“{t['大类']}”")

        lng = lat = None
        if not t["经度"] or not t["纬度"]:
            errs.append("缺坐标：请用网页“坐标拾取”补上经度、纬度")
        else:
            try:
                lng, lat = float(t["经度"]), float(t["纬度"])
            except ValueError:
                errs.append("经度、纬度必须是数字")
            else:
                if 20 < lng < 40 and 100 < lat < 140:
                    errs.append("经度、纬度好像填反了")
                elif not (YANGPU_BBOX["lng"][0] <= lng <= YANGPU_BBOX["lng"][1]
                          and YANGPU_BBOX["lat"][0] <= lat <= YANGPU_BBOX["lat"][1]):
                    errs.append(f"坐标 {lng}, {lat} 不在杨浦区范围内")

        source = sources.get(t["来源"])
        if t["来源"] and not source:
            errs.append(f"来源“{t['来源']}”不对，只能是 {' / '.join(tax['sources'])}")

        date = parse_date(v["采集日期"]) if t["采集日期"] else None
        if t["采集日期"] and not date:
            errs.append(f"采集日期“{t['采集日期']}”格式不对，请写成 2026-11-15")

        if source == "公开资料" and not t["出处"]:
            errs.append("来源是“公开资料”时，“出处”要写网址")
        refs = t["出处"].split()
        if any(not r.startswith(("http://", "https://")) for r in refs):
            errs.append("“出处”应为 http:// 或 https:// 开头的网址，多个用空格分隔")

        for head in ("说明", "设施", "骑手原话", "地址", "名称", "位置说明"):
            if PHONE_RE.search(t[head]):
                errs.append(f"“{head}”里好像有手机号，请删掉")

        key = (t["子类"], t["名称"])
        if t["名称"] and key in seen:
            errs.append(f"和第 {seen[key]} 行重复（同名同子类）")
        seen.setdefault(key, r)

        photos = []
        for name in split_list(v["照片"]):
            done = process_photo(name, errs, do_write=not args.check and not errs)
            if done:
                photos.append(done)

        if errs:
            problems.extend((r, e) for e in errs)
            continue

        if len(t["说明"]) > 120:
            warnings.append((r, "说明超过 120 字，手机上会很长，建议精简"))

        points.append({
            "id": make_id(t["名称"], sub[1]),
            "name": t["名称"],
            "category": cat["key"],
            "subtype": sub[1],
            "lng": round(lng, 6),
            "lat": round(lat, 6),
            "address": t["地址"],
            "desc": t["说明"],
            "tags": split_list(t["标签"]),
            "facilities": t["设施"],
            "hours": t["开放时间"],
            "photos": photos,
            "source": source,
            "collector": t["采集人"],
            "date": date,
            "quote": t["骑手原话"],
            "refs": refs,
            "locNote": t["位置说明"],
        })

    for r, w in warnings:
        print(f"  [提醒] 第 {r} 行：{w}")
    if problems:
        print(f"\n发现 {len(problems)} 个问题，未生成 data/points.js：")
        for r, e in problems:
            print(f"  第 {r} 行：{e}")
        sys.exit(1)
    if not points:
        print("表格里没有点位。")
        sys.exit(1)

    n_rest = sum(p["category"] == "rest" for p in points)
    print(f"校验通过：共 {len(points)} 个点位（歇脚点 {n_rest}，障碍点 {len(points) - n_rest}）。")
    if args.check:
        return

    data = {"isSample": False, "generatedAt": dt.date.today().strftime("%Y-%m-%d"), "points": points}
    body = json.dumps(data, ensure_ascii=False, indent=2)
    POINTS_JS.write_text(
        "// 由 tools/build_data.py 根据 data/points.xlsx 自动生成，请不要手动修改。\n"
        "window.RIDER_MAP_DATA = " + body + ";\n",
        encoding="utf-8",
    )
    print(f"已写入 {POINTS_JS.relative_to(ROOT)}。刷新网页即可看到。")


if __name__ == "__main__":
    main()
