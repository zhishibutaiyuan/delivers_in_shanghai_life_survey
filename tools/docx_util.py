"""生成 Word 文档的通用工具：中文字体、标题、表格、待填写占位。"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "宋体"
CN_HEAD_FONT = "黑体"
GREEN = RGBColor(0x1F, 0x8A, 0x4C)
GRAY = RGBColor(0x5D, 0x63, 0x5A)


def _set_font(style_or_run, east_asia, latin="Times New Roman"):
    font = style_or_run.font
    font.name = latin
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    _set_font(normal, CN_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(4)

    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 12)):
        st = doc.styles[name]
        _set_font(st, CN_HEAD_FONT, "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1D, 0x21, 0x1C)
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Number"):
        _set_font(doc.styles[name], CN_FONT)
    return doc


def para(doc, text="", bold=False, size=None, color=None, align=None, italic=False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align == "center":
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    return p


def rich(doc, parts, style=None):
    """parts: [(文字, {"bold": True}), "普通文字", ...]"""
    p = doc.add_paragraph(style=style)
    for part in parts:
        if isinstance(part, str):
            p.add_run(part)
        else:
            text, fmt = part
            r = p.add_run(text)
            r.bold = fmt.get("bold", False)
            if fmt.get("todo"):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def steps(doc, items):
    """手动编号的列表（每组从 1 开始）。items 中每项为 "文字" 或 ("加粗前缀", "文字")。"""
    for n, item in enumerate(items, start=1):
        prefix, text = item if isinstance(item, tuple) else (None, item)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.add_run(f"{n}. ")
        if prefix:
            p.add_run(prefix).bold = True
        p.add_run(text)


def todo(doc, text):
    """黄色高亮的“待填写”占位段落。"""
    p = doc.add_paragraph()
    r = p.add_run(f"【待填写】{text}")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def _shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def table(doc, header, rows, widths_cm, head_fill="1F8A4C", font_size=10):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, head_fill)
    # 表格跨页时重复表头
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font_size)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def kv_table(doc, pairs, widths_cm=(3, 13), font_size=10):
    """两列“字段—内容”表，左列浅灰底。"""
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    for k, v in pairs:
        cells = t.add_row().cells
        cells[0].text = ""
        r = cells[0].paragraphs[0].add_run(k)
        r.bold = True
        r.font.size = Pt(font_size)
        _shade(cells[0], "F1F3EF")
        cells[1].text = ""
        r = cells[1].paragraphs[0].add_run(v)
        r.font.size = Pt(font_size)
    for i, row in enumerate(t.rows):
        row.cells[0].width = Cm(widths_cm[0])
        row.cells[1].width = Cm(widths_cm[1])
        # 整张小表尽量不跨页：除最后一行外都“与下段同页”
        if i < len(t.rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    return t


def page_break(doc):
    doc.add_page_break()


def page_numbers(doc):
    """页脚居中页码。"""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
