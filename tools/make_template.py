"""生成发给调研组的《点位采集表.xlsx》（带下拉框和填写说明）。

用法：python tools/make_template.py
输出：docs/点位采集表.xlsx
分类改动后（data/taxonomy.js）重新运行即可。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import COLUMNS, ROOT, load_taxonomy

OUT = ROOT / "docs" / "点位采集表.xlsx"
MAX_ROWS = 300

HEAD_FILL = PatternFill("solid", fgColor="1F8A4C")
OPT_FILL = PatternFill("solid", fgColor="5D635A")
WRAP = Alignment(wrap_text=True, vertical="top")

WIDTHS = {"名称": 18, "大类": 8, "子类": 12, "经度": 12, "纬度": 11, "地址": 22, "说明": 36, "标签": 20,
          "设施": 26, "开放时间": 12, "照片": 18, "来源": 16, "采集人": 9, "采集日期": 12, "骑手原话": 30}

# 示例行：信息来自 2026 年春季实地调研记录表（新心驿站），坐标需本学期用坐标拾取补上
EXAMPLE = {
    "名称": "新心驿站（五角场街道党群服务中心）",
    "大类": "歇脚点",
    "子类": "官方驿站",
    "经度": "（到现场用网页坐标拾取）",
    "纬度": "",
    "地址": "国和路 60 号",
    "说明": "街道党群服务中心内的骑手驿站，设施齐全，可休息、饮水、充电、热饭。",
    "标签": "可饮水、可充电、卫生间、空调、微波炉、雨具",
    "设施": "免费饮用水、充电插座、休息座椅、空调、微波炉、应急药品、卫生间，可借雨具和充电宝",
    "开放时间": "（现场核实）",
    "照片": "xinxin-01.jpg",
    "来源": "实地走访",
    "采集人": "春季调研组",
    "采集日期": "2026-05-10",
    "骑手原话": "自己在路边就是舒服……非常应急的情况下才去。",
}


def main():
    tax = load_taxonomy()
    wb = Workbook()

    # ---------- 点位（要填的表） ----------
    ws = wb.active
    ws.title = "点位"
    for i, (head, _, required, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=head + (" *" if required else ""))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL if required else OPT_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = WIDTHS.get(head, 14)
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "B2"

    letter = {head: ws.cell(row=1, column=i).column_letter for i, (head, *_rest) in enumerate(COLUMNS, start=1)}

    def add_list(col, options, prompt):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="请从下拉框选择", error=prompt)
        ws.add_data_validation(dv)
        dv.add(f"{letter[col]}2:{letter[col]}{MAX_ROWS}")

    add_list("大类", [c["label"] for c in tax["categories"]], "大类只能是：歇脚点 / 障碍点")
    add_list("子类", [s["label"] for c in tax["categories"] for s in c["subtypes"]], "子类见“分类”工作表")
    add_list("来源", tax["sources"], "来源只能是：" + " / ".join(tax["sources"]))

    for col in ("经度", "纬度"):
        dv = DataValidation(type="decimal", operator="between",
                            formula1="121.4" if col == "经度" else "31.2",
                            formula2="121.7" if col == "经度" else "31.4",
                            allow_blank=True, showErrorMessage=True, errorTitle="坐标超出杨浦区范围",
                            error="经度应在 121.4~121.7，纬度应在 31.2~31.4。是不是把经度纬度填反了？")
        ws.add_data_validation(dv)
        dv.add(f"{letter[col]}2:{letter[col]}{MAX_ROWS}")
        for r in range(2, MAX_ROWS + 1):
            ws[f"{letter[col]}{r}"].number_format = "0.000000"
    for r in range(2, MAX_ROWS + 1):
        ws[f"{letter['采集日期']}{r}"].number_format = "yyyy-mm-dd"
        for col in ("说明", "设施", "骑手原话"):
            ws[f"{letter[col]}{r}"].alignment = WRAP

    # ---------- 填写说明 ----------
    guide = wb.create_sheet("填写说明")
    guide.column_dimensions["A"].width = 12
    guide.column_dimensions["B"].width = 8
    guide.column_dimensions["C"].width = 70
    guide.column_dimensions["D"].width = 48
    lines = [
        "《五角场骑手地图》点位采集表 —— 填写说明",
        "1. 每行一个点位，填在“点位”工作表。带 * 的列必填（绿色表头）。",
        "2. 点位必须是你亲自走到、或骑手当面指认的地方，不能从网上或 AI 那里抄。",
        "3. 坐标：到现场后用手机打开地图网页 → 右上角“十字”按钮（坐标拾取）→ 用我的当前位置 → 复制，粘贴到“经度”格子即可（会自动填进经度、纬度两列）。",
        "4. 照片：不拍骑手正脸、车牌、手机号；文件名用英文或拼音，如 xinxin-01.jpg，填进“照片”列，照片文件另外打包交给技术组。",
        "5. 骑手原话：只写话，不写骑手姓名；有代表性的再写，没有就空着。",
        "6. 说明写客观事实（“电动车不能进，需步行约 5 分钟”），不写评价（“保安态度差”）。",
        "",
    ]
    for r, text in enumerate(lines, start=1):
        guide.cell(row=r, column=1, value=text).font = Font(bold=(r == 1), size=13 if r == 1 else 11)

    start = len(lines) + 1
    for i, h in enumerate(["列", "必填", "怎么填", "示例（新心驿站，信息来自春季调研）"], start=1):
        c = guide.cell(row=start, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    for j, (head, _, required, hint) in enumerate(COLUMNS, start=start + 1):
        guide.cell(row=j, column=1, value=head)
        guide.cell(row=j, column=2, value="是" if required else "")
        guide.cell(row=j, column=3, value=hint).alignment = WRAP
        guide.cell(row=j, column=4, value=EXAMPLE.get(head, "")).alignment = WRAP

    # ---------- 分类 ----------
    cat = wb.create_sheet("分类")
    for i, h in enumerate(["大类", "子类", "图标", "包括哪些地方"], start=1):
        c = cat.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    r = 2
    for c in tax["categories"]:
        for s in c["subtypes"]:
            cat.cell(row=r, column=1, value=c["label"])
            cat.cell(row=r, column=2, value=s["label"])
            cat.cell(row=r, column=3, value=s["icon"])
            cat.cell(row=r, column=4, value=s["hint"])
            r += 1
    r += 1
    cat.cell(row=r, column=1, value="常用标签").font = Font(bold=True)
    cat.cell(row=r, column=2, value="、".join(tax["tags"]) + "（也可以自己写新标签）")
    for col, w in zip("ABCD", (10, 14, 6, 50)):
        cat.column_dimensions[col].width = w

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"已生成 {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
